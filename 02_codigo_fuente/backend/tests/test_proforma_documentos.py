"""
Documentos de la proforma en Word y Excel — RF-07.

Formato institucional calcado de la proforma real (CT-129-2025): datos
del oferente, concepto, notas con L/A/bloques de cada ítem, materiales en
el orden de esa proforma, condiciones, tabla de pagos y datos bancarios
con la firma — todos con `settings`, nunca hardcodeados en `documentos.py`.
El Excel es editable: subtotales y totales son fórmulas.

Dos niveles: el servicio (`app.services.documentos`) genera los bytes a
partir de una `Cotizacion` ya armada, y la API los expone como descarga
con el mismo permiso que el detalle del historial (RF-10).
"""

import io
from datetime import datetime, timezone
from decimal import Decimal
from itertools import count
from urllib.parse import quote

import pytest
from docx import Document
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.permisos import Permiso
from app.models.cotizacion import Cotizacion, ItemCotizacion, MaterialCotizado
from app.models.material import Material
from app.services import documentos
from app.services.documentos import _moneda, settings
from tests.fabricas import crear_usuario, perfil_personalizado
from tests.test_varios_inversores import _cabecera, api, h  # noqa: F401

HOY = datetime(2026, 10, 26, 15, 0, tzinfo=timezone.utc)
_secuencia = count(1)

CAMPOS_OFERENTE = [
    "HEXTRUCTURE_RAZON_SOCIAL",
    "HEXTRUCTURE_ACTIVIDAD",
    "HEXTRUCTURE_RUC",
    "HEXTRUCTURE_TELEFONO",
    "HEXTRUCTURE_CORREO",
    "HEXTRUCTURE_DIRECCION",
    "HEXTRUCTURE_REPRESENTANTE",
    "HEXTRUCTURE_CARGO_REPRESENTANTE",
    "HEXTRUCTURE_BANCO",
    "HEXTRUCTURE_BANCO_TIPO_CUENTA",
    "HEXTRUCTURE_BANCO_NUMERO_CUENTA",
    "HEXTRUCTURE_BANCO_TITULAR",
]


@pytest.fixture
def datos_oferente(monkeypatch):
    """
    Datos de prueba, con la misma forma que los reales (ver `.env`, nunca
    en el código: el repositorio es público) — alcanza para probar que
    cada campo aparece donde debe.
    """
    valores = {
        "HEXTRUCTURE_RAZON_SOCIAL": "Empresa de prueba S.A.S.",
        "HEXTRUCTURE_ACTIVIDAD": "Actividad de prueba",
        "HEXTRUCTURE_RUC": "0000000000001",
        "HEXTRUCTURE_TELEFONO": "+593 900000000",
        "HEXTRUCTURE_CORREO": "prueba@example.com",
        "HEXTRUCTURE_DIRECCION": "Dirección de prueba, Ciudad",
        "HEXTRUCTURE_REPRESENTANTE": "Representante de Prueba",
        "HEXTRUCTURE_CARGO_REPRESENTANTE": "Representante legal",
        "HEXTRUCTURE_BANCO": "Banco de Prueba",
        "HEXTRUCTURE_BANCO_TIPO_CUENTA": "Cuenta Corriente",
        "HEXTRUCTURE_BANCO_NUMERO_CUENTA": "0000000000",
        "HEXTRUCTURE_BANCO_TITULAR": "Empresa de prueba S.A.S.",
    }
    for campo, valor in valores.items():
        monkeypatch.setattr(settings, campo, valor)
    return valores


@pytest.fixture
def sin_datos_oferente(monkeypatch):
    """Nada configurado todavía (como en CI): el documento debe generarse igual."""
    for campo in CAMPOS_OFERENTE:
        monkeypatch.setattr(settings, campo, "")


def _material_o_existente(db: Session, codigo: str, nombre: str) -> Material:
    """Códigos reales y únicos en la base: si el test ya lo creó, se reusa en vez de duplicar."""
    existente = db.query(Material).filter(Material.codigo == codigo).one_or_none()
    if existente is not None:
        return existente
    material = Material(codigo=codigo, nombre=nombre, unidad="u")
    db.add(material)
    db.flush()
    return material


def _cotizacion_con_materiales(db: Session, **campos) -> Cotizacion:
    """
    Dos materiales que SÍ están en el orden de la proforma real
    (VAR1650 antes que HEX_V18MM) para probar el orden sin depender de
    sufijos únicos — por eso, a diferencia de otras fábricas de prueba,
    usa los códigos reales tal cual. Una base de pruebas por test
    (transacción revertida) hace que no choquen entre sí.
    """
    usuario = crear_usuario(db)
    hex_v18mm = _material_o_existente(db, "HEX_V18MM", "Perfil HEX 18mm")
    var1650 = _material_o_existente(db, "VAR1650", "Varilla 1650")

    cotizacion = Cotizacion(
        usuario_id=usuario.id,
        cliente_nombre=campos.pop("cliente_nombre", "SolarAndes Instalaciones"),
        cliente_empresa=campos.pop("cliente_empresa", "SolarAndes S.A."),
        cliente_email=campos.pop("cliente_email", "proyectos@solarandes.ec"),
        cliente_telefono=campos.pop("cliente_telefono", "099 123 4567"),
        proyecto_nombre=campos.pop("proyecto_nombre", "Planta El Carmen"),
        numero_proforma=campos.pop("numero_proforma", "CT-130-2026"),
        potencia_kwp=campos.pop("potencia_kwp", Decimal("45.20")),
        subtotal=Decimal("440.00"),
        addendum_porcentaje=Decimal("5.00"),
        addendum_monto=Decimal("22.00"),
        iva_porcentaje=Decimal("15.00"),
        iva_monto=Decimal("69.30"),
        total=Decimal("531.30"),
        created_at=HOY,
        items=[ItemCotizacion(paneles_largo=4, paneles_ancho=28, bloques=12)],
        materiales=[
            # A propósito en orden "equivocado": HEX_V18MM antes que
            # VAR1650, para probar que el documento los reordena.
            MaterialCotizado(material=hex_v18mm, cantidad=400, precio_unitario=Decimal("0.35")),
            MaterialCotizado(material=var1650, cantidad=120, precio_unitario=Decimal("2.50")),
        ],
        **campos,
    )
    db.add(cotizacion)
    db.flush()
    db.refresh(cotizacion)
    return cotizacion


# ─── _moneda: formato de Ecuador ──────────────────────────────────────


@pytest.mark.parametrize(
    "valor,esperado",
    [
        (Decimal("531.30"), "$ 531,30"),
        (Decimal("5025.60"), "$ 5.025,60"),
        (Decimal("1234567.89"), "$ 1.234.567,89"),
        (Decimal("0"), "$ 0,00"),
        (None, "—"),
    ],
)
def test_moneda_formato_ecuador(valor, esperado) -> None:
    assert _moneda(valor) == esperado


# ─── Orden de materiales (no alfabético) ──────────────────────────────


def test_materiales_en_el_orden_de_la_proforma_real(db: Session) -> None:
    """VAR1650 antes que HEX_V18MM, aunque se guarden al revés y H < V alfabéticamente."""
    cotizacion = _cotizacion_con_materiales(db)

    ordenados = documentos._materiales_en_orden(cotizacion)

    assert [m.material.codigo for m in ordenados] == ["VAR1650", "HEX_V18MM"]


def test_material_fuera_del_orden_conocido_cae_al_final(db: Session) -> None:
    usuario = crear_usuario(db)
    desconocido = Material(codigo="ZZZ_NUEVO", nombre="Material nuevo", unidad="u")
    var1650 = Material(codigo="VAR1650", nombre="Varilla 1650", unidad="u")
    db.add_all([desconocido, var1650])
    db.flush()
    cotizacion = Cotizacion(
        usuario_id=usuario.id,
        cliente_nombre="Cliente",
        subtotal=Decimal("1"),
        addendum_monto=Decimal("0"),
        iva_monto=Decimal("0"),
        total=Decimal("1"),
        materiales=[
            MaterialCotizado(material=desconocido, cantidad=1, precio_unitario=Decimal("1")),
            MaterialCotizado(material=var1650, cantidad=1, precio_unitario=Decimal("1")),
        ],
    )
    db.add(cotizacion)
    db.flush()

    ordenados = documentos._materiales_en_orden(cotizacion)

    assert [m.material.codigo for m in ordenados] == ["VAR1650", "ZZZ_NUEVO"]


# ─── Servicio: generar_word ────────────────────────────────────────────


def test_word_incluye_encabezado_concepto_potencia_y_cliente(db: Session, datos_oferente) -> None:
    cotizacion = _cotizacion_con_materiales(db)

    documento = Document(io.BytesIO(documentos.generar_word(cotizacion)))
    texto = "\n".join(p.text for p in documento.paragraphs)

    assert "Empresa de prueba S.A.S." in texto
    assert "Actividad de prueba" in texto
    assert "RUC: 0000000000001" in texto
    assert "Dirección de prueba, Ciudad" in texto
    assert "+593 900000000" in texto and "prueba@example.com" in texto
    assert "Proforma CT-130-2026" in texto
    assert "Concepto: Estructura de soporte reusable para la instalación en piso de módulos fotovoltaicos" in texto
    assert "Potencia: 45.20 kWp" in texto
    assert "SolarAndes Instalaciones" in texto
    assert "SolarAndes S.A." in texto
    assert "Planta El Carmen" in texto


def test_word_materiales_en_orden_con_moneda_ecuador(db: Session, datos_oferente) -> None:
    cotizacion = _cotizacion_con_materiales(db)

    documento = Document(io.BytesIO(documentos.generar_word(cotizacion)))
    filas = [[c.text for c in fila.cells] for tabla in documento.tables for fila in tabla.rows]

    assert ["Código", "Material", "Cantidad", "Precio unitario", "Subtotal"] in filas
    indice_var = next(i for i, f in enumerate(filas) if f and f[0] == "VAR1650")
    indice_hex = next(i for i, f in enumerate(filas) if f and f[0] == "HEX_V18MM")
    assert indice_var < indice_hex, "VAR1650 debe listarse antes que HEX_V18MM, como en la proforma real"
    assert ["VAR1650", "Varilla 1650 (u)", "120", "$ 2,50", "$ 300,00"] in filas
    assert ["HEX_V18MM", "Perfil HEX 18mm (u)", "400", "$ 0,35", "$ 140,00"] in filas


def test_word_totales_notas_condiciones_y_pagos(db: Session, datos_oferente) -> None:
    cotizacion = _cotizacion_con_materiales(db)

    documento = Document(io.BytesIO(documentos.generar_word(cotizacion)))
    texto = "\n".join(p.text for p in documento.paragraphs)
    filas = [[c.text for c in fila.cells] for tabla in documento.tables for fila in tabla.rows]

    assert "Subtotal: $ 440,00" in texto
    assert "Addendum (5.00%): $ 22,00" in texto
    assert "IVA (15.00%): $ 69,30" in texto
    assert "Total: $ 531,30" in texto

    # Notas fijas + configuración del ítem debajo de la primera
    assert "Las cantidades se calcularon en base al diseño de bloques configurados por el cliente." in texto
    assert "L=4 A=28 BLOQUES=12" in texto
    assert "Las estructuras se entregarán en bodegas de HEXtructure S.A.S. – Ibarra." in texto

    assert "Plazo de entrega: 20 días." in texto
    assert "Validez de la oferta: 90 días." in texto
    assert "Garantía: 1 año(s)." in texto

    assert ["Tramo", "Porcentaje", "Monto"] in filas
    assert ["Anticipo", "30%", "$ 159,39"] in filas
    assert ["Entrega", "40%", "$ 212,52"] in filas
    assert ["A 60 días de la entrega", "30%", "$ 159,39"] in filas

    assert "Banco: Banco de Prueba" in texto
    assert "Número de cuenta: 0000000000" in texto
    assert "Representante de Prueba" in texto
    assert "Representante legal" in texto


def test_word_sin_datos_del_oferente_no_falla_y_omite_las_secciones(db: Session, sin_datos_oferente) -> None:
    """Como en CI: nada configurado todavía. El documento se genera igual, sin esas líneas."""
    cotizacion = _cotizacion_con_materiales(db)

    documento = Document(io.BytesIO(documentos.generar_word(cotizacion)))
    texto = "\n".join(p.text for p in documento.paragraphs)

    assert "RUC:" not in texto
    assert "Banco" not in texto
    # Sin representante, tampoco hay bloque de firma
    assert "Representante legal" not in texto
    # El resto del documento se genera de todas formas
    assert "Proforma CT-130-2026" in texto
    assert "Total: $ 531,30" in texto


def test_word_sin_proforma_ni_datos_opcionales_del_cliente_no_falla(db: Session, sin_datos_oferente) -> None:
    """Una cotización al vuelo: sin número de proforma, empresa, correo ni teléfono."""
    cotizacion = _cotizacion_con_materiales(
        db, numero_proforma=None, cliente_empresa=None, cliente_email=None, cliente_telefono=None, proyecto_nombre=None
    )

    documento = Document(io.BytesIO(documentos.generar_word(cotizacion)))
    texto = "\n".join(p.text for p in documento.paragraphs)

    assert "Proforma (borrador)" in texto
    assert "Empresa" not in texto
    assert "Proyecto" not in texto


def test_word_sin_potencia_kwp_no_la_muestra(db: Session, sin_datos_oferente) -> None:
    cotizacion = _cotizacion_con_materiales(db, potencia_kwp=None)

    documento = Document(io.BytesIO(documentos.generar_word(cotizacion)))
    texto = "\n".join(p.text for p in documento.paragraphs)

    assert "Potencia" not in texto


# ─── Servicio: generar_excel (editable, con fórmulas) ─────────────────


def _fila_con_etiqueta(hoja, etiqueta: str, columna: str = "D") -> int:
    for fila in hoja.iter_rows():
        for celda in fila:
            if celda.coordinate.startswith(columna) and celda.value == etiqueta:
                return celda.row
    raise AssertionError(f"No se encontró la etiqueta {etiqueta!r} en la columna {columna}")


def test_excel_subtotales_y_totales_son_formulas(db: Session, datos_oferente) -> None:
    """Si Santiago cambia una cantidad, todo lo demás debe recalcularse solo."""
    cotizacion = _cotizacion_con_materiales(db)

    libro = load_workbook(io.BytesIO(documentos.generar_excel(cotizacion)), data_only=False)
    hoja = libro.active

    # Cada material: cantidad × precio, no el valor ya calculado.
    fila_hex, fila_var = (
        next(f.row for f in hoja["A"] if f.value == "HEX_V18MM"),
        next(f.row for f in hoja["A"] if f.value == "VAR1650"),
    )
    assert hoja[f"E{fila_hex}"].value == f"=C{fila_hex}*D{fila_hex}"
    assert hoja[f"E{fila_var}"].value == f"=C{fila_var}*D{fila_var}"

    fila_subtotal = _fila_con_etiqueta(hoja, "Subtotal")
    assert hoja[f"E{fila_subtotal}"].value == f"=SUM(E{min(fila_hex, fila_var)}:E{max(fila_hex, fila_var)})"

    fila_addendum = _fila_con_etiqueta(hoja, "Addendum (5.00%)")
    assert hoja[f"E{fila_addendum}"].value == f"=E{fila_subtotal}*5.00/100"

    fila_iva = _fila_con_etiqueta(hoja, "IVA (15.00%)")
    # Sobre subtotal + addendum, no solo sobre el subtotal.
    assert hoja[f"E{fila_iva}"].value == f"=(E{fila_subtotal}+E{fila_addendum})*15.00/100"

    fila_total = _fila_con_etiqueta(hoja, "Total")
    assert hoja[f"E{fila_total}"].value == f"=E{fila_subtotal}+E{fila_addendum}+E{fila_iva}"

    for etiqueta, pct in [("Anticipo", "30"), ("Entrega", "40"), ("A 60 días de la entrega", "30")]:
        fila_pago = _fila_con_etiqueta(hoja, etiqueta, columna="A")
        assert hoja[f"C{fila_pago}"].value == f"=E{fila_total}*{pct}/100", etiqueta


def test_excel_incluye_encabezado_materiales_en_orden_y_datos_bancarios(db: Session, datos_oferente) -> None:
    cotizacion = _cotizacion_con_materiales(db)

    libro = load_workbook(io.BytesIO(documentos.generar_excel(cotizacion)))
    filas = [list(fila) for fila in libro.active.iter_rows(values_only=True) if any(v is not None for v in fila)]
    aplanado = [v for fila in filas for v in fila if isinstance(v, str)]

    assert "Empresa de prueba S.A.S." in aplanado
    assert "RUC: 0000000000001" in aplanado
    assert ["Código", "Material", "Cantidad", "Precio unitario", "Subtotal"] in filas

    indice_var = next(i for i, f in enumerate(filas) if f[0] == "VAR1650")
    indice_hex = next(i for i, f in enumerate(filas) if f[0] == "HEX_V18MM")
    assert indice_var < indice_hex

    assert "Banco" in aplanado and "Banco de Prueba" in aplanado
    assert "Representante de Prueba" in aplanado


def test_excel_sin_datos_del_oferente_no_falla(db: Session, sin_datos_oferente) -> None:
    cotizacion = _cotizacion_con_materiales(db)

    libro = load_workbook(io.BytesIO(documentos.generar_excel(cotizacion)))
    aplanado = [v for fila in libro.active.iter_rows(values_only=True) for v in fila if isinstance(v, str)]

    assert not any(v.startswith("RUC:") for v in aplanado)
    assert "Banco" not in aplanado


# ─── nombre_archivo ─────────────────────────────────────────────────────


def test_nombre_archivo_usa_la_proforma_o_el_id_si_es_borrador(db: Session) -> None:
    con_proforma = _cotizacion_con_materiales(db, numero_proforma="CT-130-2026")
    assert documentos.nombre_archivo(con_proforma, "docx") == "Proforma CT-130-2026.docx"

    borrador = _cotizacion_con_materiales(db, numero_proforma=None)
    assert documentos.nombre_archivo(borrador, "xlsx") == f"Proforma cotizacion-{borrador.id}.xlsx"


# ─── API: descarga ───────────────────────────────────────────────────

TIPOS_MEDIO = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


@pytest.mark.parametrize("extension", ["docx", "xlsx"])
def test_descarga_responde_el_archivo_con_su_tipo_y_nombre(api, db, h, extension) -> None:  # noqa: F811
    cotizacion = _cotizacion_con_materiales(db)

    r = api.get(f"/api/cotizaciones/{cotizacion.id}/proforma.{extension}", headers=h)

    assert r.status_code == 200
    assert r.headers["content-type"] == TIPOS_MEDIO[extension]
    nombre_codificado = quote(f"Proforma {cotizacion.numero_proforma}.{extension}")
    assert nombre_codificado in r.headers["content-disposition"]
    assert len(r.content) > 0


@pytest.mark.parametrize("extension", ["docx", "xlsx"])
def test_descarga_404_si_no_existe_la_cotizacion(api, h, extension) -> None:  # noqa: F811
    assert api.get(f"/api/cotizaciones/999999/proforma.{extension}", headers=h).status_code == 404


@pytest.mark.parametrize("extension", ["docx", "xlsx"])
def test_descarga_exige_el_mismo_permiso_que_el_detalle(api, db, extension) -> None:  # noqa: F811
    cotizacion = _cotizacion_con_materiales(db)

    con_historial = _cabecera(crear_usuario(db, perfil=perfil_personalizado(db, Permiso.HISTORIAL)))
    assert api.get(f"/api/cotizaciones/{cotizacion.id}/proforma.{extension}", headers=con_historial).status_code == 200

    sin_permiso = _cabecera(crear_usuario(db, perfil=perfil_personalizado(db, Permiso.CLIENTES)))
    assert api.get(f"/api/cotizaciones/{cotizacion.id}/proforma.{extension}", headers=sin_permiso).status_code == 403
