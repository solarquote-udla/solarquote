"""
Documentos de la proforma en Word y Excel — RF-07.

Formato institucional, calcado de la proforma real de HEXtructure
(CT-129-2025, proyecto EMIHANA, 10/01/2026): encabezado del oferente,
concepto, potencia en kWp si la cotización la tiene, notas con la
configuración L/A/bloques de cada ítem, materiales en el orden de esa
proforma (no alfabético), condiciones comerciales, tabla de pagos y
datos bancarios con la firma del representante legal.

El Word es un documento fijo. El Excel es editable: los subtotales por
material, el subtotal general, el addendum, el IVA, el total y la tabla
de pagos son fórmulas, no valores — si Santiago cambia una cantidad, todo
lo demás se recalcula solo.

El dinero se maneja en `Decimal` y se imprime con el formato de Ecuador
($ 5.025,60), el mismo de la proforma real — distinto a como se ve en la
interfaz, a propósito (ver `utils/dinero.ts` en el frontend).
"""

import io
from decimal import Decimal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.core.fechas import ZONA_ECUADOR
from app.models.cotizacion import Cotizacion, MaterialCotizado
from app.services.cotizacion import CODIGO_LOGISTICA

# Orden de la proforma real (CT-129-2025), no alfabético. Un material que
# no esté aquí (fuera del catálogo de las fórmulas f(L,A,B) y de LOG) cae
# al final, ordenado por código, para no romper si el catálogo crece.
ORDEN_MATERIALES_PROFORMA = (
    "VAR1650",
    "VAR1350",
    "HEX_V18MM",
    "SOP_A",
    "SOP_V",
    "SOP_V_E",
    "ANT_B",
    "ANT_VAR",
    "CLAMP_F",
    "CLAMP_I",
    CODIGO_LOGISTICA,
)


def _orden_material(material: MaterialCotizado) -> tuple[int, str]:
    codigo = material.material.codigo
    try:
        return (ORDEN_MATERIALES_PROFORMA.index(codigo), "")
    except ValueError:
        return (len(ORDEN_MATERIALES_PROFORMA), codigo)


def _materiales_en_orden(cotizacion: Cotizacion) -> list[MaterialCotizado]:
    return sorted(cotizacion.materiales, key=_orden_material)


def _fecha(cotizacion: Cotizacion) -> str:
    return cotizacion.created_at.astimezone(ZONA_ECUADOR).strftime("%d/%m/%Y")


def _moneda(valor: Decimal | None) -> str:
    """Formato de Ecuador: punto de millares, coma decimal — "$ 5.025,60"."""
    if valor is None:
        return "—"
    valor = valor.quantize(Decimal("0.01"))
    negativo = valor < 0
    entero, _, decimales = f"{abs(valor):.2f}".partition(".")
    grupos = []
    while len(entero) > 3:
        grupos.insert(0, entero[-3:])
        entero = entero[:-3]
    grupos.insert(0, entero)
    signo = "-" if negativo else ""
    return f"{signo}$ {'.'.join(grupos)},{decimales}"


def _subtotal_material(material: MaterialCotizado) -> Decimal:
    return Decimal(material.cantidad) * material.precio_unitario


def _notas(cotizacion: Cotizacion) -> list[str]:
    """
    Notas fijas de la proforma, con la configuración L/A/bloques de cada
    ítem debajo de la primera ("las cantidades se calcularon en base
    a..."), tal como en la proforma real.
    """
    notas = list(settings.PROFORMA_NOTAS)
    configuraciones = [
        f"L={item.paneles_largo} A={item.paneles_ancho} BLOQUES={item.bloques}" for item in cotizacion.items
    ]
    if not notas:
        return configuraciones
    return [notas[0], *configuraciones, *notas[1:]]


def _pagos(total: Decimal | None) -> list[tuple[str, Decimal, Decimal | None]]:
    """(etiqueta, porcentaje, monto) para cada tramo de la forma de pago."""
    tramos = [
        ("Anticipo", settings.PROFORMA_PAGO_ANTICIPO_PCT),
        ("Entrega", settings.PROFORMA_PAGO_ENTREGA_PCT),
        ("A 60 días de la entrega", settings.PROFORMA_PAGO_60_DIAS_PCT),
    ]
    return [
        (etiqueta, pct, (total * pct / Decimal("100")) if total is not None else None) for etiqueta, pct in tramos
    ]


def nombre_archivo(cotizacion: Cotizacion, extension: str) -> str:
    base = cotizacion.numero_proforma or f"cotizacion-{cotizacion.id}"
    return f"Proforma {base}.{extension}"


# ─── Word ───────────────────────────────────────────────────────────


def generar_word(cotizacion: Cotizacion) -> bytes:
    documento = Document()

    titulo = documento.add_heading(settings.HEXTRUCTURE_RAZON_SOCIAL or "HEXtructure S.A.S.", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    datos_oferente = [
        settings.HEXTRUCTURE_ACTIVIDAD,
        f"RUC: {settings.HEXTRUCTURE_RUC}" if settings.HEXTRUCTURE_RUC else None,
        settings.HEXTRUCTURE_DIRECCION or None,
        " · ".join(p for p in (settings.HEXTRUCTURE_TELEFONO, settings.HEXTRUCTURE_CORREO) if p) or None,
    ]
    for linea in datos_oferente:
        if linea:
            parrafo = documento.add_paragraph(linea)
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitulo = documento.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corrida = subtitulo.add_run(
        f"Proforma {cotizacion.numero_proforma}" if cotizacion.numero_proforma else "Proforma (borrador)"
    )
    corrida.bold = True
    corrida.font.size = Pt(14)
    documento.add_paragraph(f"Fecha: {_fecha(cotizacion)}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    documento.add_paragraph(f"Concepto: {settings.PROFORMA_CONCEPTO}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cotizacion.potencia_kwp is not None:
        documento.add_paragraph(f"Potencia: {cotizacion.potencia_kwp} kWp").alignment = WD_ALIGN_PARAGRAPH.CENTER

    documento.add_heading("Cliente", level=2)
    datos_cliente = [("Nombre", cotizacion.cliente_nombre)]
    if cotizacion.cliente_empresa:
        datos_cliente.append(("Empresa", cotizacion.cliente_empresa))
    if cotizacion.cliente_email:
        datos_cliente.append(("Correo", cotizacion.cliente_email))
    if cotizacion.cliente_telefono:
        datos_cliente.append(("Teléfono", cotizacion.cliente_telefono))
    if cotizacion.proyecto_nombre:
        datos_cliente.append(("Proyecto", cotizacion.proyecto_nombre))
    for etiqueta, valor in datos_cliente:
        documento.add_paragraph(f"{etiqueta}: {valor}")

    documento.add_heading("Materiales", level=2)
    tabla = documento.add_table(rows=1, cols=5)
    tabla.style = "Light Grid Accent 1"
    encabezados = ["Código", "Material", "Cantidad", "Precio unitario", "Subtotal"]
    for celda, texto in zip(tabla.rows[0].cells, encabezados):
        celda.text = texto
        for parrafo in celda.paragraphs:
            for corrida in parrafo.runs:
                corrida.bold = True

    for material in _materiales_en_orden(cotizacion):
        fila = tabla.add_row().cells
        valores = [
            material.material.codigo,
            f"{material.material.nombre} ({material.material.unidad})",
            str(material.cantidad),
            _moneda(material.precio_unitario),
            _moneda(_subtotal_material(material)),
        ]
        for celda, texto in zip(fila, valores):
            celda.text = texto

    documento.add_heading("Notas", level=2)
    for nota in _notas(cotizacion):
        documento.add_paragraph(nota, style="List Bullet")

    documento.add_heading("Totales", level=2)
    filas_totales = [
        ("Subtotal", cotizacion.subtotal),
        (f"Addendum ({cotizacion.addendum_porcentaje}%)", cotizacion.addendum_monto),
        (f"IVA ({cotizacion.iva_porcentaje}%)" if cotizacion.iva_porcentaje is not None else "IVA", cotizacion.iva_monto),
        ("Total", cotizacion.total),
    ]
    for etiqueta, monto in filas_totales:
        parrafo = documento.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        corrida = parrafo.add_run(f"{etiqueta}: {_moneda(monto)}")
        if etiqueta == "Total":
            corrida.bold = True

    documento.add_heading("Condiciones", level=2)
    documento.add_paragraph(f"Plazo de entrega: {settings.PROFORMA_PLAZO_DIAS} días.")
    documento.add_paragraph(f"Validez de la oferta: {settings.PROFORMA_VALIDEZ_DIAS} días.")
    documento.add_paragraph(f"Garantía: {settings.PROFORMA_GARANTIA_ANIOS} año(s).")

    documento.add_heading("Forma de pago", level=2)
    tabla_pagos = documento.add_table(rows=1, cols=3)
    tabla_pagos.style = "Light Grid Accent 1"
    for celda, texto in zip(tabla_pagos.rows[0].cells, ["Tramo", "Porcentaje", "Monto"]):
        celda.text = texto
        for parrafo in celda.paragraphs:
            for corrida in parrafo.runs:
                corrida.bold = True
    for etiqueta, pct, monto in _pagos(cotizacion.total):
        fila = tabla_pagos.add_row().cells
        for celda, texto in zip(fila, [etiqueta, f"{pct}%", _moneda(monto)]):
            celda.text = texto

    if settings.HEXTRUCTURE_BANCO:
        documento.add_heading("Datos bancarios", level=2)
        documento.add_paragraph(f"Banco: {settings.HEXTRUCTURE_BANCO}")
        if settings.HEXTRUCTURE_BANCO_TIPO_CUENTA:
            documento.add_paragraph(f"Tipo de cuenta: {settings.HEXTRUCTURE_BANCO_TIPO_CUENTA}")
        if settings.HEXTRUCTURE_BANCO_NUMERO_CUENTA:
            documento.add_paragraph(f"Número de cuenta: {settings.HEXTRUCTURE_BANCO_NUMERO_CUENTA}")
        if settings.HEXTRUCTURE_BANCO_TITULAR:
            documento.add_paragraph(f"Titular: {settings.HEXTRUCTURE_BANCO_TITULAR}")

    if settings.HEXTRUCTURE_REPRESENTANTE:
        documento.add_paragraph()
        firma = documento.add_paragraph(settings.HEXTRUCTURE_REPRESENTANTE)
        firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if settings.HEXTRUCTURE_CARGO_REPRESENTANTE:
            cargo = documento.add_paragraph(settings.HEXTRUCTURE_CARGO_REPRESENTANTE)
            cargo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    salida = io.BytesIO()
    documento.save(salida)
    return salida.getvalue()


# ─── Excel ──────────────────────────────────────────────────────────

FORMATO_MONEDA_EXCEL = '"$" #,##0.00'


def generar_excel(cotizacion: Cotizacion) -> bytes:
    """
    Igual que el Word, pero editable: los subtotales, el subtotal
    general, el addendum, el IVA y el total son fórmulas sobre las
    cantidades y precios de las filas, no valores fijos.
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Proforma"
    negrita = Font(bold=True)

    hoja["A1"] = settings.HEXTRUCTURE_RAZON_SOCIAL or "HEXtructure S.A.S."
    hoja["A1"].font = Font(bold=True, size=14)
    fila = 2
    for linea in (
        settings.HEXTRUCTURE_ACTIVIDAD,
        f"RUC: {settings.HEXTRUCTURE_RUC}" if settings.HEXTRUCTURE_RUC else None,
        settings.HEXTRUCTURE_DIRECCION or None,
        " · ".join(p for p in (settings.HEXTRUCTURE_TELEFONO, settings.HEXTRUCTURE_CORREO) if p) or None,
    ):
        if linea:
            hoja.cell(row=fila, column=1, value=linea)
            fila += 1

    fila += 1
    hoja.cell(row=fila, column=1, value=f"Proforma {cotizacion.numero_proforma}" if cotizacion.numero_proforma else "Proforma (borrador)").font = negrita
    fila += 1
    hoja.cell(row=fila, column=1, value=f"Fecha: {_fecha(cotizacion)}")
    fila += 1
    hoja.cell(row=fila, column=1, value=f"Concepto: {settings.PROFORMA_CONCEPTO}")
    fila += 1
    if cotizacion.potencia_kwp is not None:
        hoja.cell(row=fila, column=1, value=f"Potencia: {cotizacion.potencia_kwp} kWp")
        fila += 1

    fila += 1
    datos_cliente = [("Cliente", cotizacion.cliente_nombre)]
    if cotizacion.cliente_empresa:
        datos_cliente.append(("Empresa", cotizacion.cliente_empresa))
    if cotizacion.cliente_email:
        datos_cliente.append(("Correo", cotizacion.cliente_email))
    if cotizacion.cliente_telefono:
        datos_cliente.append(("Teléfono", cotizacion.cliente_telefono))
    if cotizacion.proyecto_nombre:
        datos_cliente.append(("Proyecto", cotizacion.proyecto_nombre))
    for etiqueta, valor in datos_cliente:
        hoja.cell(row=fila, column=1, value=etiqueta).font = negrita
        hoja.cell(row=fila, column=2, value=valor)
        fila += 1

    fila += 1
    encabezados = ["Código", "Material", "Cantidad", "Precio unitario", "Subtotal"]
    for columna, texto in enumerate(encabezados, start=1):
        hoja.cell(row=fila, column=columna, value=texto).font = negrita
    fila += 1

    primera_fila_materiales = fila
    for material in _materiales_en_orden(cotizacion):
        hoja.cell(row=fila, column=1, value=material.material.codigo)
        hoja.cell(row=fila, column=2, value=f"{material.material.nombre} ({material.material.unidad})")
        hoja.cell(row=fila, column=3, value=material.cantidad)
        # openpyxl guarda Decimal tal cual (XLSX no tiene un tipo numérico
        # propio: al abrirlo vuelve a ser float, pero no se pierde precisión
        # aquí, en el cálculo).
        hoja.cell(row=fila, column=4, value=material.precio_unitario).number_format = FORMATO_MONEDA_EXCEL
        # Editable: cantidad × precio, no el valor ya calculado.
        celda_subtotal = hoja.cell(row=fila, column=5, value=f"=C{fila}*D{fila}")
        celda_subtotal.number_format = FORMATO_MONEDA_EXCEL
        fila += 1
    ultima_fila_materiales = fila - 1

    fila += 1
    fila_subtotal = fila
    hoja.cell(row=fila, column=4, value="Subtotal").alignment = Alignment(horizontal="right")
    celda_subtotal_general = hoja.cell(row=fila, column=5, value=f"=SUM(E{primera_fila_materiales}:E{ultima_fila_materiales})")
    celda_subtotal_general.number_format = FORMATO_MONEDA_EXCEL
    fila += 1

    fila_addendum = fila
    hoja.cell(row=fila, column=4, value=f"Addendum ({cotizacion.addendum_porcentaje}%)").alignment = Alignment(horizontal="right")
    celda_addendum = hoja.cell(
        row=fila, column=5, value=f"=E{fila_subtotal}*{cotizacion.addendum_porcentaje}/100"
    )
    celda_addendum.number_format = FORMATO_MONEDA_EXCEL
    fila += 1

    fila_iva = fila
    etiqueta_iva = f"IVA ({cotizacion.iva_porcentaje}%)" if cotizacion.iva_porcentaje is not None else "IVA"
    hoja.cell(row=fila, column=4, value=etiqueta_iva).alignment = Alignment(horizontal="right")
    # Sobre el subtotal más el addendum, no solo el subtotal.
    celda_iva = hoja.cell(
        row=fila,
        column=5,
        value=f"=(E{fila_subtotal}+E{fila_addendum})*{cotizacion.iva_porcentaje or 0}/100",
    )
    celda_iva.number_format = FORMATO_MONEDA_EXCEL
    fila += 1

    fila_total = fila
    celda_etiqueta_total = hoja.cell(row=fila, column=4, value="Total")
    celda_etiqueta_total.alignment = Alignment(horizontal="right")
    celda_etiqueta_total.font = negrita
    celda_total = hoja.cell(row=fila, column=5, value=f"=E{fila_subtotal}+E{fila_addendum}+E{fila_iva}")
    celda_total.number_format = FORMATO_MONEDA_EXCEL
    celda_total.font = negrita
    fila += 2

    hoja.cell(row=fila, column=1, value="Notas").font = negrita
    fila += 1
    for nota in _notas(cotizacion):
        hoja.cell(row=fila, column=1, value=f"• {nota}")
        fila += 1

    fila += 1
    hoja.cell(row=fila, column=1, value=f"Plazo de entrega: {settings.PROFORMA_PLAZO_DIAS} días.")
    fila += 1
    hoja.cell(row=fila, column=1, value=f"Validez de la oferta: {settings.PROFORMA_VALIDEZ_DIAS} días.")
    fila += 1
    hoja.cell(row=fila, column=1, value=f"Garantía: {settings.PROFORMA_GARANTIA_ANIOS} año(s).")
    fila += 2

    hoja.cell(row=fila, column=1, value="Forma de pago").font = negrita
    fila += 1
    for columna, texto in enumerate(["Tramo", "Porcentaje", "Monto"], start=1):
        hoja.cell(row=fila, column=columna, value=texto).font = negrita
    fila += 1
    for etiqueta, pct, _ in _pagos(cotizacion.total):
        hoja.cell(row=fila, column=1, value=etiqueta)
        hoja.cell(row=fila, column=2, value=f"{pct}%")
        # También fórmula, sobre la celda del total: cambia si cambia una cantidad.
        celda_pago = hoja.cell(row=fila, column=3, value=f"=E{fila_total}*{pct}/100")
        celda_pago.number_format = FORMATO_MONEDA_EXCEL
        fila += 1

    if settings.HEXTRUCTURE_BANCO:
        fila += 1
        hoja.cell(row=fila, column=1, value="Datos bancarios").font = negrita
        fila += 1
        for etiqueta, valor in [
            ("Banco", settings.HEXTRUCTURE_BANCO),
            ("Tipo de cuenta", settings.HEXTRUCTURE_BANCO_TIPO_CUENTA),
            ("Número de cuenta", settings.HEXTRUCTURE_BANCO_NUMERO_CUENTA),
            ("Titular", settings.HEXTRUCTURE_BANCO_TITULAR),
        ]:
            if valor:
                hoja.cell(row=fila, column=1, value=etiqueta).font = negrita
                hoja.cell(row=fila, column=2, value=valor)
                fila += 1

    if settings.HEXTRUCTURE_REPRESENTANTE:
        fila += 1
        hoja.cell(row=fila, column=1, value=settings.HEXTRUCTURE_REPRESENTANTE)
        fila += 1
        if settings.HEXTRUCTURE_CARGO_REPRESENTANTE:
            hoja.cell(row=fila, column=1, value=settings.HEXTRUCTURE_CARGO_REPRESENTANTE)

    for columna, ancho in zip(range(1, 6), [16, 36, 12, 16, 14]):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
